# -*- coding: utf-8 -*-
"""سازنده فایل Word رساله سطح ۳ — مطابق شیوه‌نامه حوزوی (فونت B Lotus / B Titr)."""
import re
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = "/home/user/Hadi/thesis_src"
OUT = "/home/user/Hadi/Resaleh_Sath3_Hadi_Shabestani_Final2.docx"

TITLE_MAIN = "عدول ولی فقیه از تعهدات بین‌المللی دولت اسلامی"
TITLE_SUB = "از منظر فقه امامیه و حقوق بین‌الملل"
AUTHOR = "هادی شبستانی"

footnotes = []  # متن پاورقی‌ها به ترتیب


# ---------- ابزارهای پایه ----------

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        for key, val in kwargs.items():
            element.set(qn(f"w:{key}"), val)


def style_run(run, font="B Lotus", size=14, bold=False, color=None, italic=False):
    run.font.name = font
    r = run._r
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), font)
    sz = rPr.find(qn("w:sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        rPr.append(sz)
    sz.set(qn("w:val"), str(int(size * 2)))
    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = OxmlElement("w:szCs")
        rPr.append(szCs)
    szCs.set(qn("w:val"), str(int(size * 2)))
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, before=0, after=6, line=1.4):
    pPr = p._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    bidi.set(qn("w:val"), "1")
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.widow_control = True


def set_para_ltr(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, line=1.4):
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.widow_control = True


def is_latin(text):
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return False
    return sum(1 for c in letters if ord(c) < 128) / len(letters) > 0.5


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|\{\{fn:.*?\}\})", re.DOTALL)


def add_footnote_ref(paragraph, note_text, font="B Lotus", size=14):
    footnotes.append(note_text)
    fid = len(footnotes) + 1  # شناسه‌ها از ۲ شروع می‌شوند
    run = paragraph.add_run()
    style_run(run, font=font, size=size)
    rPr = run._r.get_or_add_rPr()
    va = OxmlElement("w:vertAlign")
    va.set(qn("w:val"), "superscript")
    rPr.append(va)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), str(fid))
    run._r.append(ref)


def add_runs(paragraph, text, font="B Lotus", size=14, bold=False, color=None, rtl=False):
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            style_run(run, font=font, size=size, bold=bold, color=color)
            if rtl:
                rPr = run._r.get_or_add_rPr()
                rPr.append(OxmlElement("w:rtl"))
        tok = m.group(0)
        if tok.startswith("{{fn:"):
            add_footnote_ref(paragraph, tok[5:-2].strip(), font=font, size=size)
        else:
            run = paragraph.add_run(tok[2:-2])
            style_run(run, font=font, size=size, bold=True, color=color)
            if rtl:
                rPr = run._r.get_or_add_rPr()
                rPr.append(OxmlElement("w:rtl"))
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        style_run(run, font=font, size=size, bold=bold, color=color)
        if rtl:
            rPr = run._r.get_or_add_rPr()
            rPr.append(OxmlElement("w:rtl"))


def add_field(paragraph, instr, font="B Lotus", size=10, bold=False, color=None):
    r1 = paragraph.add_run()
    style_run(r1, font=font, size=size, bold=bold, color=color)
    c1 = OxmlElement("w:fldChar")
    c1.set(qn("w:fldCharType"), "begin")
    r1._r.append(c1)
    r2 = paragraph.add_run()
    style_run(r2, font=font, size=size, bold=bold, color=color)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    style_run(r3, font=font, size=size, bold=bold, color=color)
    c3 = OxmlElement("w:fldChar")
    c3.set(qn("w:fldCharType"), "end")
    r3._r.append(c3)


def set_outline(style, level):
    pPr = style.element.get_or_add_pPr()
    ol = pPr.find(qn("w:outlineLvl"))
    if ol is None:
        ol = OxmlElement("w:outlineLvl")
        pPr.append(ol)
    ol.set(qn("w:val"), str(level))


def add_bottom_border(paragraph, color="1F3864", sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def add_top_border(paragraph, color="BFBFBF", sz="6"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), sz)
    top.set(qn("w:space"), "6")
    top.set(qn("w:color"), color)
    pbdr.append(top)


def add_frame_borders(paragraph, color="1F3864", sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for edge in ("top", "bottom"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "8")
        el.set(qn("w:color"), color)
        pbdr.append(el)
    pPr.append(pbdr)


def add_box_border(paragraph, color="BFBFBF", sz="6"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "8")
        el.set(qn("w:color"), color)
        pbdr.append(el)
    pPr.append(pbdr)


def add_shading(paragraph, fill="F2F2F2"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def set_keep_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    kn = pPr.find(qn("w:keepNext"))
    if kn is None:
        kn = OxmlElement("w:keepNext")
        pPr.append(kn)
    kn.set(qn("w:val"), "1")


def add_rule(doc, indent=3, after=12, color="1F3864", sz="12"):
    p = doc.add_paragraph()
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=after)
    p.paragraph_format.right_indent = Cm(indent)
    p.paragraph_format.left_indent = Cm(indent)
    add_bottom_border(p, color=color, sz=sz)
    return p


def add_page_border(section, color="1F3864", sz="12"):
    sectPr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "24")
        el.set(qn("w:color"), color)
        borders.append(el)
    sectPr.append(borders)


# ---------- جدول ----------

def add_table(doc, rows, header=True):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = 1  # center
    tblPr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tblPr.append(bidi)
    for i, row in enumerate(rows):
        is_head = header and i == 0
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=2, before=2)
            add_runs(p, text.strip(), font="B Lotus", size=12, bold=is_head)
            if is_head:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "D9D9D9")
                tcPr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def parse_table_block(lines):
    rows = []
    for ln in lines:
        if re.match(r"^\|\s*[-:]+\s*(\|\s*[-:]+\s*)*\|?\s*$", ln):
            continue
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


# ---------- بلوک‌های محتوا ----------

first_content_block = True


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_bakhsh(doc, title):
    global first_content_block
    if not first_content_block:
        page_break(doc)
    first_content_block = False
    for _ in range(3):
        doc.add_paragraph()
    add_rule(doc, indent=4, after=12)
    p = doc.add_paragraph(style="TH-Bakhsh")
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=12)
    add_runs(p, title, font="B Titr", size=32, bold=True, rtl=True)
    tag_bookmark(p)
    add_rule(doc, indent=4, after=10)
    page_break(doc)


def add_fasl(doc, title):
    global first_content_block
    if not first_content_block:
        page_break(doc)
    first_content_block = False
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph(style="TH-Fasl")
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=12, before=6)
    add_frame_borders(p)
    set_keep_next(p)
    add_runs(p, title, font="B Titr", size=22, bold=True, rtl=True)
    tag_bookmark(p)


def add_heading(doc, title, style, font, size):
    p = doc.add_paragraph(style=style)
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=8, before=14)
    set_keep_next(p)
    if style == "TH-Goftar":
        add_shading(p, "EDEDED")
    add_runs(p, title, font=font, size=size, bold=True, rtl=True)
    if style in ("TH-Goftar", "TH-Band"):
        tag_bookmark(p)


def add_body(doc, text, indent=True):
    p = doc.add_paragraph(style="TH-Body")
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)
    add_runs(p, text, font="B Lotus", size=14)


def add_quote(doc, text, full_width=False):
    p = doc.add_paragraph(style="TH-Quote")
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=12, before=12)
    pf = p.paragraph_format
    if not full_width:
        pf.right_indent = Cm(1)
        pf.left_indent = Cm(1)
    add_shading(p, "F2F2F2")
    add_box_border(p)
    add_runs(p, text, font="B Badr", size=13, bold=True)


def add_bullet(doc, text, full_width=False):
    p = doc.add_paragraph(style="TH-Body")
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=6)
    if not full_width:
        p.paragraph_format.right_indent = Cm(0.75)
    run = p.add_run("• ")
    style_run(run, font="B Lotus", size=14, bold=True)
    add_runs(p, text, font="B Lotus", size=14)


def build_content_file(doc, path):
    with open(path, encoding="utf-8") as f:
        lines = [ln.rstrip("\n") for ln in f]
    no_indent = True
    i = 0
    while i < len(lines):
        ln = lines[i].strip()
        if not ln:
            i += 1
            continue
        if ln.startswith("%%"):
            add_bakhsh(doc, ln[2:].strip())
            no_indent = True
        elif ln.startswith("#### "):
            add_heading(doc, ln[5:].strip(), "TH-Sub", "B Lotus", 14)
            no_indent = True
        elif ln.startswith("### "):
            add_heading(doc, ln[4:].strip(), "TH-Band", "B Traffic", 14)
            no_indent = True
        elif ln.startswith("## "):
            add_heading(doc, ln[3:].strip(), "TH-Goftar", "B Titr", 16)
            no_indent = True
        elif ln.startswith("# "):
            add_fasl(doc, ln[2:].strip())
            no_indent = True
        elif ln == ">":
            pass
        elif ln.startswith("> "):
            add_quote(doc, ln[2:].strip(), full_width=True)
            no_indent = True
        elif ln.startswith("- "):
            add_bullet(doc, ln[2:].strip(), full_width=True)
            no_indent = True
        elif ln.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip())
                i += 1
            add_table(doc, parse_table_block(tbl))
            no_indent = True
            continue
        else:
            add_body(doc, ln, indent=not no_indent)
            no_indent = False
        i += 1


# ---------- صفحات مقدماتی ----------

def add_cover(doc):
    for _ in range(2):
        doc.add_paragraph()
    items = [
        ("حوزه‌های علمیه", "B Titr", 16, False, 6),
        ("رساله علمی سطح سه", "B Titr", 22, True, 12),
        ("__RULE__", None, 0, False, 12),
        (TITLE_MAIN, "B Titr", 22, True, 6),
        (TITLE_SUB, "B Titr", 17, True, 12),
        ("__RULE__", None, 0, False, 12),
        ("نگارنده: " + AUTHOR, "B Lotus", 16, True, 10),
        ("استاد راهنما: …………………………", "B Lotus", 14, False, 6),
        ("استاد مشاور: …………………………", "B Lotus", 14, False, 6),
        ("گرایش: فقه و اصول", "B Lotus", 14, False, 12),
        ("سال تحصیلی ۱۴۰۵ ـ ۱۴۰۶", "B Lotus", 14, False, 6),
    ]
    for text, font, size, bold, after in items:
        p = doc.add_paragraph()
        set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=after)
        if text == "__RULE__":
            p.paragraph_format.right_indent = Cm(3)
            p.paragraph_format.left_indent = Cm(3)
            add_bottom_border(p)
        elif text:
            add_runs(p, text, font=font, size=size, bold=bold)


def add_basmalah(doc):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=12)
    add_runs(p, "بِسْمِ اللَّهِ الرَّحْمَنِ الرَّحِيمِ", font="B Titr", size=22, bold=True)
    add_rule(doc, indent=5, after=10)
    page_break(doc)


FRONT_TOC_TITLES = ("چکیده", "پیشگفتار")
FRONT_CENTER_TITLES = ("تقدیم", "سپاسگزاری")


def front_heading(doc, title, new_page=True, framed=False):
    if new_page:
        page_break(doc)
    if title in FRONT_TOC_TITLES:
        p = doc.add_paragraph(style="TH-Front")
    else:
        p = doc.add_paragraph()
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=12, before=6)
    set_keep_next(p)
    if framed:
        add_frame_borders(p)
    add_runs(p, title, font="B Titr", size=18, bold=True, rtl=True)
    if title in FRONT_TOC_TITLES:
        tag_bookmark(p)


def build_front_file(doc, path):
    with open(path, encoding="utf-8") as f:
        lines = [ln.rstrip("\n") for ln in f]
    centered, first, no_indent = False, True, True
    i = 0
    while i < len(lines):
        ln = lines[i].strip()
        if not ln:
            i += 1
            continue
        if ln.startswith("%%"):
            i += 1
            continue
        if ln == ">":
            i += 1
            continue
        if ln.startswith("# "):
            title = ln[2:].strip()
            front_heading(doc, title, new_page=not first)
            centered = (title in FRONT_CENTER_TITLES) or first
            if first:
                add_rule(doc, indent=3, after=12)
            first = False
            no_indent = True
        elif ln.startswith("> "):
            add_quote(doc, ln[2:].strip())
            first = False
            no_indent = True
        elif ln.startswith("- "):
            add_bullet(doc, ln[2:].strip())
            first = False
            no_indent = True
        elif ln.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip())
                i += 1
            add_table(doc, parse_table_block(tbl))
            first = False
            no_indent = True
            continue
        elif ln.startswith("کلیدواژه‌ها"):
            p = doc.add_paragraph()
            set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=10, before=12)
            run = p.add_run("کلیدواژه‌ها: ")
            style_run(run, font="B Lotus", size=14, bold=True)
            rest = ln.split(":", 1)[-1].strip()
            add_runs(p, rest, font="B Lotus", size=14)
            first = False
            no_indent = False
        elif is_latin(ln):
            p = doc.add_paragraph()
            set_para_ltr(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=10)
            if not no_indent:
                p.paragraph_format.first_line_indent = Cm(0.5)
            add_runs(p, ln, font="B Lotus", size=13)
            first = False
            no_indent = False
        else:
            p = doc.add_paragraph()
            if centered:
                set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)
            else:
                set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)
                if not no_indent:
                    p.paragraph_format.first_line_indent = Cm(0.5)
            add_runs(p, ln, font="B Lotus", size=14)
            first = False
            no_indent = False
        i += 1


# ---------- فهرست مطالب کتابی: عنوان + نقطه‌چین + شماره صفحه ----------

TOC_MARKS = []  # (سطح، عنوان) به ترتیب سند

def scan_toc_marks():
    TOC_MARKS.clear()
    with open(f"{SRC}/00_front.md", encoding="utf-8") as _f:
        for _ln in _f:
            _s = _ln.strip()
            if _s.startswith("# ") and _s[2:].strip() in FRONT_TOC_TITLES:
                TOC_MARKS.append((2, _s[2:].strip()))
    for _fn in ["01_moqaddame.md", "02_bakhsh1_fasl1.md", "03_bakhsh1_fasl2.md",
                "04_bakhsh2_fasl3.md", "05_bakhsh2_fasl4.md", "06_bakhsh3_fasl5.md",
                "07_natije.md", "08_manabe.md", "09_payvast.md"]:
        with open(f"{SRC}/{_fn}", encoding="utf-8") as _f:
            for _ln in _f:
                _s = _ln.strip()
                if _s.startswith("%% ") and _fn != "00_front.md":
                    TOC_MARKS.append((1, _s[3:].strip()))
                elif _s.startswith("#### "):
                    continue
                elif _s.startswith("### "):
                    TOC_MARKS.append((4, _s[4:].strip()))
                elif _s.startswith("## "):
                    TOC_MARKS.append((3, _s[3:].strip()))
                elif _s.startswith("# "):
                    TOC_MARKS.append((2, _s[2:].strip()))


def toc_mark_name(i):
    return f"_Toc{i + 1:06d}"


_mark_seq = [0]

def tag_bookmark(paragraph):
    _mark_seq[0] += 1
    name = toc_mark_name(_mark_seq[0] - 1)
    bs = OxmlElement("w:bookmarkStart")
    bs.set(qn("w:id"), str(_mark_seq[0]))
    bs.set(qn("w:name"), name)
    be = OxmlElement("w:bookmarkEnd")
    be.set(qn("w:id"), str(_mark_seq[0]))
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(bs)
    else:
        paragraph._p.insert(0, bs)
    paragraph._p.append(be)


def add_pageref(paragraph, name, font="B Lotus", size=13):
    r1 = paragraph.add_run()
    style_run(r1, font=font, size=size)
    c1 = OxmlElement("w:fldChar")
    c1.set(qn("w:fldCharType"), "begin")
    r1._r.append(c1)
    r2 = paragraph.add_run()
    style_run(r2, font=font, size=size)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" PAGEREF {name} \\h "
    r2._r.append(it)
    r3 = paragraph.add_run()
    style_run(r3, font=font, size=size)
    c3 = OxmlElement("w:fldChar")
    c3.set(qn("w:fldCharType"), "separate")
    r3._r.append(c3)
    r5 = paragraph.add_run()
    style_run(r5, font=font, size=size)
    c5 = OxmlElement("w:fldChar")
    c5.set(qn("w:fldCharType"), "end")
    r5._r.append(c5)


def add_toc(doc):
    front_heading(doc, "فهرست مطالب", new_page=True, framed=True)
    p = doc.add_paragraph()
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)
    # فیلد فهرست خودکار
    r1 = p.add_run()
    style_run(r1, font="B Lotus", size=14)
    c1 = OxmlElement("w:fldChar")
    c1.set(qn("w:fldCharType"), "begin")
    r1._r.append(c1)
    r2 = p.add_run()
    style_run(r2, font="B Lotus", size=14)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "1-4" \\h \\z \\u '
    r2._r.append(it)
    r3 = p.add_run()
    style_run(r3, font="B Lotus", size=14)
    c3 = OxmlElement("w:fldChar")
    c3.set(qn("w:fldCharType"), "separate")
    r3._r.append(c3)
    # نتیجه ایستای فهرست به سبک کتاب: عنوان + نقطه‌چین + شماره صفحه (فیلد زنده)
    for _i, (_lv, _tx) in enumerate(TOC_MARKS):
        _ep = doc.add_paragraph(style=f"TOC{_lv}")
        _epPr = _ep._p.get_or_add_pPr()
        _bidi = OxmlElement("w:bidi")
        _bidi.set(qn("w:val"), "1")
        _epPr.append(_bidi)
        _jc = OxmlElement("w:jc")
        _jc.set(qn("w:val"), "right")
        _epPr.append(_jc)
        add_runs(_ep, _tx, font="B Lotus", size=13, rtl=True)
        _tr = _ep.add_run("\t")
        style_run(_tr, font="B Lotus", size=13)
        add_pageref(_ep, toc_mark_name(_i))
    pe = doc.add_paragraph()
    set_para_rtl(pe, align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)
    r5 = pe.add_run()
    style_run(r5, font="B Lotus", size=14)
    c5 = OxmlElement("w:fldChar")
    c5.set(qn("w:fldCharType"), "end")
    r5._r.append(c5)
    hint = doc.add_paragraph()
    set_para_rtl(hint, align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)
    add_runs(hint, "راهنما: شماره صفحه‌ها با باز شدن فایل در Word خودکار کامل می‌شود؛ اگر نشد، روی فهرست کلیک راست کرده و «Update Field» و سپس «Update entire table» را بزنید؛ سپس این جمله را حذف کنید.",
             font="B Lotus", size=11, color=RGBColor(0x88, 0x88, 0x88))


# ---------- استایل‌ها و بخش‌ها ----------

def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "B Lotus"
    normal.font.size = Pt(14)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), "B Lotus")
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set(qn("w:val"), "28")
    pPr = normal.element.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)

    styles = doc.styles
    for name, base, outline in [
        ("TH-Bakhsh", "Normal", 0),
        ("TH-Fasl", "Normal", 1),
        ("TH-Front", "Normal", 1),
        ("TH-Goftar", "Normal", 2),
        ("TH-Band", "Normal", 3),
        ("TH-Sub", "Normal", None),
        ("TH-Body", "Normal", None),
        ("TH-Quote", "Normal", None),
    ]:
        if name in styles:
            st = styles[name]
        else:
            st = styles.add_style(name, 1)  # paragraph
            st.base_style = styles[base]
        if outline is not None:
            set_outline(st, outline)

    # استایل‌های فهرست مطالب (راست‌به‌چپ + نقطه‌چین شماره صفحه + تورفتگی پلکانی)
    styles_el = doc.styles.element
    toc_indent = {1: 0, 2: 284, 3: 567, 4: 851}
    for idx in range(1, 5):
        sid = f"TOC{idx}"
        if sid in [s.style_id for s in styles]:
            continue
        st = OxmlElement("w:style")
        st.set(qn("w:type"), "paragraph")
        st.set(qn("w:styleId"), sid)
        nm = OxmlElement("w:name")
        nm.set(qn("w:val"), f"toc {idx}")
        st.append(nm)
        based = OxmlElement("w:basedOn")
        based.set(qn("w:val"), "Normal")
        st.append(based)
        pPr = OxmlElement("w:pPr")
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        pPr.append(bidi)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "right")
        pPr.append(jc)
        if toc_indent[idx]:
            ind = OxmlElement("w:ind")
            ind.set(qn("w:right"), str(toc_indent[idx]))
            pPr.append(ind)
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "left")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), "9072")
        tabs.append(tab)
        pPr.append(tabs)
        st.append(pPr)
        rPr = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(attr), "B Lotus")
        rPr.append(rf)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "26")
        rPr.append(sz)
        szCs = OxmlElement("w:szCs")
        szCs.set(qn("w:val"), "26")
        rPr.append(szCs)
        st.append(rPr)
        styles_el.append(st)


def setup_page(section):
    sectPr = section._sectPr
    for pb in sectPr.findall(qn("w:pgBorders")):
        sectPr.remove(pb)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.right_margin = Cm(3)
    section.left_margin = Cm(2)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)


def setup_main_header(section):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ""
    set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, after=2, before=0)
    add_bottom_border(p, color="BFBFBF", sz="6")
    add_field(p, 'STYLEREF "TH-Fasl" \\* MERGEFORMAT', font="B Lotus", size=10)


def setup_main_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.2
    add_top_border(p)
    run = p.add_run("صفحه ")
    style_run(run, font="B Lotus", size=11, bold=True)
    add_field(p, "PAGE \\* MERGEFORMAT", font="B Lotus", size=11, bold=True)


# ---------- پاورقی (تزریق پس از ذخیره) ----------

def build_footnotes_xml():
    parts = ['<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">']
    parts.append('<w:footnote w:id="0" w:type="separator"><w:p><w:pPr><w:bidi w:val="1"/><w:jc w:val="right"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>')
    parts.append('<w:footnote w:id="1" w:type="continuationSeparator"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>')
    for i, text in enumerate(footnotes, start=2):
        parts.append(
            f'<w:footnote w:id="{i}"><w:p><w:pPr><w:bidi w:val="1"/><w:jc w:val="right"/>'
            f'<w:spacing w:after="40" w:line="240" w:lineRule="auto"/></w:pPr>'
            f'<w:r><w:rPr><w:rFonts w:ascii="B Lotus" w:hAnsi="B Lotus" w:cs="B Lotus"/>'
            f'<w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr><w:footnoteRef/></w:r>'
            f'<w:r><w:rPr><w:rFonts w:ascii="B Lotus" w:hAnsi="B Lotus" w:cs="B Lotus"/>'
            f'<w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr>'
            f'<w:t xml:space="preserve"> </w:t></w:r>'
            f'<w:r><w:rPr><w:rFonts w:ascii="B Lotus" w:hAnsi="B Lotus" w:cs="B Lotus"/>'
            f'<w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr>'
            f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'
            f'</w:p></w:footnote>'
        )
    parts.append("</w:footnotes>")
    return "".join(parts).encode("utf-8")


def inject_footnotes(docx_path):
    with zipfile.ZipFile(docx_path, "r") as zin:
        items = {name: zin.read(name) for name in zin.namelist()}
    items["word/footnotes.xml"] = build_footnotes_xml()
    # رابطه در document.xml.rels
    rels = items["word/_rels/document.xml.rels"].decode("utf-8")
    ids = [int(m) for m in re.findall(r'Id="rId(\d+)"', rels)]
    new_id = (max(ids) if ids else 0) + 100
    anchor = "</Relationships>"
    add = (f'<Relationship Id="rId{new_id}" '
           f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
           f'Target="footnotes.xml"/>' + anchor)
    items["word/_rels/document.xml.rels"] = rels.replace(anchor, add).encode("utf-8")
    # نوع محتوا
    ct = items["[Content_Types].xml"].decode("utf-8")
    anchor2 = "</Types>"
    add2 = ('<Override PartName="/word/footnotes.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>' + anchor2)
    items["[Content_Types].xml"] = ct.replace(anchor2, add2).encode("utf-8")
    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items.items():
            zout.writestr(name, data)


# ---------- اجرای اصلی ----------

def main():
    doc = Document()
    setup_styles(doc)

    # بخش ۱: جلد (با قاب دور صفحه)
    cover_section = doc.sections[0]
    setup_page(cover_section)
    add_page_border(cover_section)
    cp = doc.core_properties
    cp.title = TITLE_MAIN + " " + TITLE_SUB
    cp.author = AUTHOR
    cp.subject = "رساله علمی سطح سه حوزه علمیه"

    # به‌روزرسانی خودکار فیلدها هنگام بازشدن
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    doc.settings.element.append(uf)

    add_cover(doc)

    # بخش ۲: صفحات آغازین
    doc.add_section(WD_SECTION.NEW_PAGE)
    front_section = doc.sections[1]
    setup_page(front_section)
    add_basmalah(doc)
    scan_toc_marks()
    build_front_file(doc, f"{SRC}/00_front.md")
    add_toc(doc)

    # بخش ۳: متن اصلی؛ شماره صفحه از ۱ + سرصفحه هوشمند
    doc.add_section(WD_SECTION.NEW_PAGE)
    main_section = doc.sections[2]
    setup_page(main_section)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    main_section._sectPr.append(pg)
    setup_main_header(main_section)
    setup_main_footer(main_section)
    for fn in ["01_moqaddame.md", "02_bakhsh1_fasl1.md", "03_bakhsh1_fasl2.md",
               "04_bakhsh2_fasl3.md", "05_bakhsh2_fasl4.md", "06_bakhsh3_fasl5.md",
               "07_natije.md", "08_manabe.md"]:
        build_content_file(doc, f"{SRC}/{fn}")
    for fn in ["09_payvast.md"]:
        build_content_file(doc, f"{SRC}/{fn}")

    doc.save(OUT)
    inject_footnotes(OUT)

    # آمار
    import glob
    words = 0
    for fp in glob.glob(f"{SRC}/*.md"):
        with open(fp, encoding="utf-8") as f:
            words += len(f.read().split())
    print(f"OK: {OUT}")
    print(f"words(source)={words} footnotes={len(footnotes)} toc-marks={len(TOC_MARKS)} bookmarks={_mark_seq[0]}")
    if _mark_seq[0] != len(TOC_MARKS):
        print("WARNING: bookmark/mark mismatch!")


if __name__ == "__main__":
    main()
